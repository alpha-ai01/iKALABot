import os
import csv
from pypdf import PdfReader
from docx import Document
import openpyxl

def process_document(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    
    # Text-based source/doc formats
    text_formats = {
        '.txt', '.md', '.py', '.js', '.mjs', '.ts', '.tsx', '.css', '.html', '.htm',
        '.json', '.xml', '.yaml', '.yml'
    }
    
    if ext in text_formats:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
            
    elif ext == '.pdf':
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
        
    elif ext == '.docx':
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs])
        
    elif ext == '.csv':
        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            return "\n".join([",".join(row) for row in reader])
            
    elif ext == '.xlsx':
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text = ""
        for sheet in wb.worksheets:
            text += f"Sheet: {sheet.title}\n"
            for row in sheet.iter_rows(values_only=True):
                text += ",".join([str(cell) for cell in row if cell is not None]) + "\n"
        return text
        
    else:
        return "รูปแบบไฟล์ไม่รองรับ"
